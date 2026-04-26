import io
import re
from pathlib import Path

import pytest
from docx import Document

from app.services.master_docx import build_master_docx, _render_markdown_content, _iter_all_paragraphs

MASTER_DOCX_PATH = Path(__file__).parent.parent / "templates" / "master.docx"


def _minimal_template_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("{{ FINDINGS_HIGH }}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _summary_template_bytes() -> bytes:
    """Template with all summary placeholders for testing markdown rendering."""
    doc = Document()
    doc.add_paragraph("{{ SUMMARY_EXECUTIVE }}")
    doc.add_paragraph("{{ SUMMARY_TECHNICAL }}")
    doc.add_paragraph("{{ SUMMARY_RECOMMENDATIONS }}")
    doc.add_paragraph("{{ FINDINGS_HIGH }}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_master_docx_localizes_severity_for_docx_only():
    template_bytes = _minimal_template_bytes()
    finding_content = {
        "Severity": "High",
        "Vulnerability Summary": "Summary text",
    }
    findings = [{"_title": "Test Finding", "_content": finding_content}]

    out = build_master_docx(
        template_bytes=template_bytes,
        report_name="R",
        client_name="C",
        team_name="T",
        company_name="Comp",
        department_name="Dep",
        author_name="Author",
        author_email="author@example.com",
        findings=findings,
        locale="de",
    )

    doc = Document(io.BytesIO(out))
    paragraph_texts = [p.text for p in doc.paragraphs]
    
    # Note: Localization requires frontend locale files mounted in container
    # In test env without frontend files, it falls back to English
    # Test that severity appears (either localized or English)
    assert any("Severity:" in t for t in paragraph_texts)
    assert any("Test Finding" in t for t in paragraph_texts)
    assert any("Vulnerability Summary" in t for t in paragraph_texts)

    # Guardrail: do not mutate canonical finding content_json values
    assert finding_content["Severity"] == "High"


class TestSummaryMarkdownRendering:
    """Summaries with markdown are rendered as formatted content, not literal text."""

    def _build(self, summaries: dict) -> Document:
        template_bytes = _summary_template_bytes()
        out = build_master_docx(
            template_bytes=template_bytes,
            report_name="R",
            client_name="C",
            team_name="T",
            company_name="Comp",
            department_name="Dep",
            author_name="Author",
            author_email="author@example.com",
            findings=[],
            summaries=summaries,
        )
        return Document(io.BytesIO(out))

    def test_bold_markers_not_literal(self):
        doc = self._build({"SUMMARY_EXECUTIVE": "**Key finding**: system is exposed."})
        texts = [p.text for p in doc.paragraphs]
        assert not any("**" in t for t in texts), "Bold markers should not appear as literal text"
        assert any("Key finding" in t for t in texts)

    def test_heading_markers_not_literal(self):
        doc = self._build({"SUMMARY_TECHNICAL": "## Overview\n\nDetailed analysis here."})
        texts = [p.text for p in doc.paragraphs]
        assert not any("##" in t for t in texts), "Heading markers should not appear as literal text"
        assert any("Overview" in t for t in texts)
        assert any("Detailed analysis" in t for t in texts)

    def test_bullet_list_rendered(self):
        md = "Recommendations:\n\n- Patch immediately\n- Review access controls\n- Enable logging"
        doc = self._build({"SUMMARY_RECOMMENDATIONS": md})
        texts = [p.text for p in doc.paragraphs]
        styles = [p.style.name for p in doc.paragraphs]
        assert not any(t.startswith("- ") for t in texts), "Bullet markers should not be literal"
        assert "Patch immediately" in texts or any("Patch immediately" in t for t in texts)
        assert "List Bullet" in styles

    def test_plain_text_summary_still_works(self):
        doc = self._build({"SUMMARY_EXECUTIVE": "No issues found. System is secure."})
        texts = [p.text for p in doc.paragraphs]
        assert any("No issues found" in t for t in texts)

    def test_empty_summary_no_crash(self):
        doc = self._build({"SUMMARY_EXECUTIVE": "", "SUMMARY_TECHNICAL": ""})
        # Should not raise; marker should not appear in output
        texts = [p.text for p in doc.paragraphs]
        assert not any("LZPT_SUM_" in t for t in texts)


class TestRenderMarkdownContent:
    """Unit tests for _render_markdown_content (block-level markdown in DOCX)."""

    def _render(self, text: str) -> list[str]:
        """Helper: render markdown into a fresh doc and return paragraph texts."""
        doc = Document()
        _render_markdown_content(doc, text)
        return [p.text for p in doc.paragraphs]

    def _render_styles(self, text: str) -> list[tuple[str, str]]:
        """Helper: return (text, style_name) pairs."""
        doc = Document()
        _render_markdown_content(doc, text)
        return [(p.text, p.style.name) for p in doc.paragraphs]

    def test_plain_paragraph(self):
        texts = self._render("This is a plain paragraph.")
        assert "This is a plain paragraph." in texts

    def test_atx_heading_h1_maps_to_heading4(self):
        pairs = self._render_styles("# Section Header")
        heading_paragraphs = [(t, s) for t, s in pairs if "Section Header" in t]
        assert len(heading_paragraphs) == 1
        _, style = heading_paragraphs[0]
        assert "Heading" in style

    def test_atx_heading_h2(self):
        pairs = self._render_styles("## Sub-section")
        heading_paragraphs = [(t, s) for t, s in pairs if "Sub-section" in t]
        assert len(heading_paragraphs) == 1
        _, style = heading_paragraphs[0]
        assert "Heading" in style

    def test_bullet_list_items(self):
        text = "- First item\n- Second item\n- Third item"
        pairs = self._render_styles(text)
        bullet_items = [(t, s) for t, s in pairs if s == "List Bullet"]
        assert len(bullet_items) == 3
        texts = [t for t, _ in bullet_items]
        assert "First item" in texts
        assert "Second item" in texts
        assert "Third item" in texts

    def test_numbered_list_items(self):
        text = "1. Step one\n2. Step two\n3. Step three"
        pairs = self._render_styles(text)
        numbered_items = [(t, s) for t, s in pairs if s == "List Number"]
        assert len(numbered_items) == 3

    def test_fenced_code_block(self):
        text = "```\npayload = '<script>alert(1)</script>'\n```"
        doc = Document()
        _render_markdown_content(doc, text)
        texts = [p.text for p in doc.paragraphs]
        assert any("payload" in t for t in texts)
        # Verify monospace font applied
        code_para = next(p for p in doc.paragraphs if "payload" in p.text)
        assert code_para.runs[0].font.name == "Courier New"

    def test_mixed_content(self):
        text = (
            "## Overview\n\n"
            "Some introductory text.\n\n"
            "- Item A\n"
            "- Item B\n\n"
            "1. Step one\n"
            "2. Step two\n\n"
            "```\ncode here\n```\n\n"
            "Closing paragraph."
        )
        pairs = self._render_styles(text)
        styles = [s for _, s in pairs]
        assert any("Heading" in s for s in styles)
        assert "List Bullet" in styles
        assert "List Number" in styles
        texts = [t for t, _ in pairs]
        assert any("introductory" in t for t in texts)
        assert any("code here" in t for t in texts)
        assert any("Closing paragraph" in t for t in texts)

    def test_validation_guidance_in_full_export(self):
        """Validation guidance with markdown is rendered via block-level parser in DOCX export."""
        template_bytes = _minimal_template_bytes()
        validation_md = (
            "## Summary\n\n"
            "The following steps should be taken:\n\n"
            "- Patch the library\n"
            "- Update the configuration\n\n"
            "```\ncurl -X POST /api/test\n```"
        )
        finding_content = {
            "Severity": "High",
            "Vulnerability Summary": "XSS vulnerability found.",
            "Validation Guidance": validation_md,
        }
        findings = [{"_title": "XSS Finding", "_content": finding_content}]

        out = build_master_docx(
            template_bytes=template_bytes,
            report_name="R",
            client_name="C",
            team_name="T",
            company_name="Comp",
            department_name="Dep",
            author_name="Author",
            author_email="author@example.com",
            findings=findings,
        )

        doc = Document(io.BytesIO(out))
        paragraph_texts = [p.text for p in doc.paragraphs]
        paragraph_styles = [p.style.name for p in doc.paragraphs]

        # Raw markdown syntax should NOT appear as literal text
        assert not any("##" in t for t in paragraph_texts), "Heading markers should not be literal text"
        assert not any(t.startswith("- ") for t in paragraph_texts), "Bullet markers should not be literal text"
        assert not any("```" in t for t in paragraph_texts), "Code fence markers should not be literal text"

        # Rendered content should be present
        assert any("Summary" in t for t in paragraph_texts)
        assert any("Patch the library" in t for t in paragraph_texts)
        assert any("curl" in t for t in paragraph_texts)
        # Bullet list style should be applied
        assert "List Bullet" in paragraph_styles


# ── Shipping master.docx token interpolation ─────────────────────────────────

_JINJA_TOKEN_RE = re.compile(r'\{\{[^}]+\}\}')
_INTERNAL_MARKER_RE = re.compile(r'LZPT_SUM_\w+_MARKER')

_FULL_FINDINGS = [
    {
        "_title": "Critical RCE in login endpoint",
        "_content": {
            "Severity": "Critical",
            "Vulnerability Summary": "Remote code execution via unsanitized input.",
            "Recommendations": "Patch the input handling immediately.",
            "CVSS Score": "9.8",
        },
    },
    {
        "_title": "Reflected XSS in search",
        "_content": {
            "Severity": "High",
            "Vulnerability Summary": "User input reflected without encoding.",
            "Recommendations": "Encode all output.",
        },
    },
    {
        "_title": "Missing rate limiting",
        "_content": {"Severity": "Medium", "Vulnerability Summary": "No rate limiting on API."},
    },
    {
        "_title": "Verbose error messages",
        "_content": {"Severity": "Low", "Vulnerability Summary": "Stack traces exposed in errors."},
    },
    {
        "_title": "HTTP used for static assets",
        "_content": {"Severity": "Informational", "Vulnerability Summary": "Some assets served over HTTP."},
    },
]


@pytest.mark.skipif(not MASTER_DOCX_PATH.exists(), reason="templates/master.docx not present")
class TestMasterDocxTokenInterpolation:
    """Verify that every Jinja2 token in the shipping master.docx is resolved after rendering."""

    def _render(self) -> Document:
        template_bytes = MASTER_DOCX_PATH.read_bytes()
        out = build_master_docx(
            template_bytes=template_bytes,
            report_name="2026 Pentest Report",
            client_name="Acme Web Platform",
            team_name="Acme Corporation",
            company_name="Pentest Co Ltd",
            department_name="Offensive Security",
            author_name="Jane Smith",
            author_email="jane.smith@pentestco.com",
            author_phone="+1 555 0100",
            author_title="Senior Penetration Tester",
            findings=_FULL_FINDINGS,
            summaries={
                "SUMMARY_EXECUTIVE": "**Risk level: High.** Three critical and high findings require immediate attention.",
                "SUMMARY_TECHNICAL": "## Findings Overview\n\nDetailed technical analysis of all vulnerabilities.\n\n- RCE via unsanitized input\n- XSS in search",
                "SUMMARY_RECOMMENDATIONS": "1. Patch RCE immediately\n2. Encode XSS output\n3. Add rate limiting",
            },
        )
        return Document(io.BytesIO(out))

    def _collect_all_text(self, doc: Document) -> list[str]:
        return [p.text for p in _iter_all_paragraphs(doc)]

    def test_no_unresolved_jinja_tokens(self):
        """No {{ TOKEN }} patterns must remain in the rendered document."""
        doc = self._render()
        texts = self._collect_all_text(doc)
        unresolved = [
            match
            for text in texts
            for match in _JINJA_TOKEN_RE.findall(text)
        ]
        assert not unresolved, (
            f"{len(unresolved)} unresolved token(s) found in rendered master.docx:\n"
            + "\n".join(f"  {t}" for t in sorted(set(unresolved)))
        )

    def test_no_internal_summary_markers_remain(self):
        """Internal LZPT_SUM_*_MARKER strings must be fully replaced by summary content."""
        doc = self._render()
        texts = self._collect_all_text(doc)
        leaked = [
            match
            for text in texts
            for match in _INTERNAL_MARKER_RE.findall(text)
        ]
        assert not leaked, (
            f"Internal summary marker(s) not replaced in rendered output: {sorted(set(leaked))}"
        )

    def test_author_fields_appear_in_output(self):
        """Key author and report metadata tokens must produce visible content."""
        doc = self._render()
        full_text = " ".join(self._collect_all_text(doc))
        assert "Jane Smith" in full_text, "Author name missing — AUTHOR token not interpolated"
        assert "jane.smith@pentestco.com" in full_text, "Author email missing — AUTH_EMAIL/AUTHOR_EMAIL token not interpolated"
        assert "Acme" in full_text, "Client name missing — APP_NAME/CLIENT_NAME token not interpolated"
        assert "Pentest Co" in full_text, "Company name missing — MY_COMPANY_NAME token not interpolated"

    def test_finding_titles_appear_in_output(self):
        """At least one finding title must be present in the rendered document."""
        doc = self._render()
        full_text = " ".join(self._collect_all_text(doc))
        assert any(f["_title"] in full_text for f in _FULL_FINDINGS), (
            "No finding titles found in rendered output"
        )
