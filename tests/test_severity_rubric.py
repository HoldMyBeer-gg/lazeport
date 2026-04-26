"""
Unit tests for severity rubric extraction and enforcement.

All test data uses synthetic/fake DOCX content only. No real assessment data.
"""

import io
import pytest
from docx import Document

from app.services.severity_rubric import (
    extract_rubric_from_docx_bytes,
    SeverityRubric,
    enforce_severity_label,
    SEVERITY_SYNONYMS,
)


def create_test_docx_with_table() -> bytes:
    """Create a fake DOCX with severity table."""
    doc = Document()
    doc.add_heading("Security Assessment Report", level=1)
    doc.add_heading("Severity Ratings", level=2)
    
    table = doc.add_table(rows=6, cols=2)
    table.rows[0].cells[0].text = "Severity Level"
    table.rows[0].cells[1].text = "Definition"
    
    table.rows[1].cells[0].text = "Critical"
    table.rows[1].cells[1].text = "System compromise, immediate action required"
    
    table.rows[2].cells[0].text = "High"
    table.rows[2].cells[1].text = "Significant security risk"
    
    table.rows[3].cells[0].text = "Medium"
    table.rows[3].cells[1].text = "Moderate security concern"
    
    table.rows[4].cells[0].text = "Low"
    table.rows[4].cells[1].text = "Minor security issue"
    
    table.rows[5].cells[0].text = "Informational"
    table.rows[5].cells[1].text = "Security-relevant information"
    
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_test_docx_with_headings() -> bytes:
    """Create a fake DOCX with severity definitions under headings."""
    doc = Document()
    doc.add_heading("Severity Definitions", level=2)
    
    doc.add_paragraph("Critical: Complete system compromise")
    doc.add_paragraph("High: Significant data exposure")
    doc.add_paragraph("Medium: Limited security impact")
    doc.add_paragraph("Low: Minimal risk")
    doc.add_paragraph("Informational: No direct security impact")
    
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_test_docx_empty() -> bytes:
    """Create a fake DOCX with no severity information."""
    doc = Document()
    doc.add_heading("Test Report", level=1)
    doc.add_paragraph("This document has no severity rubric.")
    
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestRubricExtraction:
    """Test severity rubric extraction from DOCX templates."""
    
    def test_extract_from_table(self):
        """Test extraction from table with Severity headers."""
        docx_bytes = create_test_docx_with_table()
        rubric = extract_rubric_from_docx_bytes("test", docx_bytes)
        
        assert rubric.source == "template"
        assert "Critical" in rubric.allowed_labels
        assert "High" in rubric.allowed_labels
        assert "Medium" in rubric.allowed_labels
        assert "Low" in rubric.allowed_labels
        assert "Informational" in rubric.allowed_labels
        assert len(rubric.definitions) >= 5
        assert "System compromise" in rubric.definitions.get("Critical", "")
    
    def test_extract_from_headings(self):
        """Test extraction from paragraph definitions."""
        docx_bytes = create_test_docx_with_headings()
        rubric = extract_rubric_from_docx_bytes("test", docx_bytes)
        
        assert rubric.source == "template"
        assert "Critical" in rubric.allowed_labels
        assert "Complete system compromise" in rubric.definitions.get("Critical", "")
    
    def test_extract_fallback_to_default(self):
        """Test fallback to defaults when no rubric found."""
        docx_bytes = create_test_docx_empty()
        rubric = extract_rubric_from_docx_bytes("test", docx_bytes)
        
        assert rubric.source == "default"
        assert rubric.allowed_labels == ["Critical", "High", "Medium", "Low", "Informational"]
        assert rubric.definitions == {}
    
    def test_extraction_caching(self):
        """Test that extraction results are cached."""
        docx_bytes = create_test_docx_with_table()
        
        # First call
        rubric1 = extract_rubric_from_docx_bytes("test", docx_bytes)
        
        # Second call with same content should return cached result
        rubric2 = extract_rubric_from_docx_bytes("test", docx_bytes)
        
        assert rubric1.allowed_labels == rubric2.allowed_labels
        assert rubric1.source == rubric2.source


class TestSeverityRubric:
    """Test SeverityRubric methods."""
    
    def test_is_valid_label(self):
        """Test valid label checking."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low"],
            definitions={},
            source="test"
        )
        
        assert rubric.is_valid_label("Critical")
        assert rubric.is_valid_label("critical")  # Case-insensitive
        assert rubric.is_valid_label("HIGH")
        assert not rubric.is_valid_label("Informational")  # Not in allowed list
        assert not rubric.is_valid_label("Severe")
    
    def test_clamp_severity_valid_label(self):
        """Test that valid labels pass through unchanged."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="test"
        )
        
        clamped, reason = rubric.clamp_severity("High")
        assert clamped == "High"
        assert reason is None
    
    def test_clamp_severity_synonym_map(self):
        """Test that synonyms are mapped deterministically."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="test"
        )
        
        # Test synonyms
        clamped, reason = rubric.clamp_severity("Severe")
        assert clamped == "Critical"
        assert "synonym_map" in reason
        
        clamped, reason = rubric.clamp_severity("Major")
        assert clamped == "High"
        assert "synonym_map" in reason
        
        clamped, reason = rubric.clamp_severity("Moderate")
        assert clamped == "Medium"
        assert "synonym_map" in reason
        
        clamped, reason = rubric.clamp_severity("Information")
        assert clamped == "Informational"
        assert "synonym_map" in reason
    
    def test_clamp_severity_invalid_to_lowest(self):
        """Test that invalid labels clamp to lowest allowed label."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="test"
        )
        
        clamped, reason = rubric.clamp_severity("Super Critical")
        assert clamped == "Informational"  # Lowest
        assert "invalid_label" in reason
        
        clamped, reason = rubric.clamp_severity("Unknown")
        assert clamped == "Informational"
        assert "clamped_to_lowest" in reason
    
    def test_clamp_severity_empty_label(self):
        """Test that empty labels clamp to lowest."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="test"
        )
        
        clamped, reason = rubric.clamp_severity("")
        assert clamped == "Informational"
        assert "empty_label" in reason


class TestEnforceSeverityLabel:
    """Test severity label enforcement in structured output."""
    
    def test_enforce_valid_severity(self):
        """Test that valid severity passes through."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="template"
        )
        
        structured = {"Severity": "High", "Summary": "Test finding"}
        result = enforce_severity_label(structured, rubric)
        
        assert result["Severity"] == "High"
        assert "_severity_flags" in result
        assert "rubric_source=template" in result["_severity_flags"]
    
    def test_enforce_clamps_invalid_severity(self):
        """Test that invalid severity is clamped."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="template"
        )
        
        structured = {"Severity": "Super High", "Summary": "Test finding"}
        result = enforce_severity_label(structured, rubric)
        
        assert result["Severity"] == "Informational"  # Clamped to lowest
        assert "_severity_flags" in result
        assert any("invalid_label" in flag for flag in result["_severity_flags"])
    
    def test_enforce_uses_synonym_map(self):
        """Test that synonyms are properly mapped."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="template"
        )
        
        structured = {"Severity": "Severe", "Summary": "Test finding"}
        result = enforce_severity_label(structured, rubric)
        
        assert result["Severity"] == "Critical"
        assert "_severity_flags" in result
        assert any("synonym_map" in flag for flag in result["_severity_flags"])
    
    def test_enforce_adds_missing_severity(self):
        """Test that missing severity field is added."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="default"
        )
        
        structured = {"Summary": "Test finding"}
        result = enforce_severity_label(structured, rubric)
        
        assert "Severity" in result
        assert result["Severity"] == "Informational"  # Lowest
        assert "_severity_flags" in result
        assert any("missing_severity" in flag for flag in result["_severity_flags"])
    
    def test_enforce_case_insensitive(self):
        """Test that enforcement is case-insensitive."""
        rubric = SeverityRubric(
            allowed_labels=["Critical", "High", "Medium", "Low", "Informational"],
            definitions={},
            source="template"
        )
        
        structured = {"severity": "high", "Summary": "Test finding"}  # Lowercase
        result = enforce_severity_label(structured, rubric)
        
        assert result["severity"] == "High"  # Capitalized
        assert "_severity_flags" in result


class TestSynonymMap:
    """Test that the synonym map is complete and correct."""
    
    def test_all_synonyms_map_to_valid_defaults(self):
        """Test that all synonyms map to valid default severity labels."""
        for synonym, canonical in SEVERITY_SYNONYMS.items():
            assert canonical in ["Critical", "High", "Medium", "Low", "Informational"]
    
    def test_no_ambiguous_synonyms(self):
        """Test that no synonym maps to multiple values."""
        seen = {}
        for synonym, canonical in SEVERITY_SYNONYMS.items():
            assert synonym.lower() not in seen, f"Duplicate synonym: {synonym}"
            seen[synonym.lower()] = canonical


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
