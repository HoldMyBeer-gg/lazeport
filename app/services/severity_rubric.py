"""
Severity rubric extraction and enforcement.

Provides deterministic severity label validation using the master template
as the single source of truth when available.
"""

import io
import logging
import re
import hashlib
from typing import Dict, List, Optional, Tuple
from dataclasses import dataclass

from docx import Document

log = logging.getLogger(__name__)

# Default severity labels (fallback when template has no rubric)
DEFAULT_SEVERITY_LABELS = ["Critical", "High", "Medium", "Low", "Informational"]

# Fixed synonym map for deterministic clamping (no "intent inference")
SEVERITY_SYNONYMS = {
    # Critical synonyms
    "severe": "Critical",
    "urgent": "Critical",
    "catastrophic": "Critical",
    
    # High synonyms
    "major": "High",
    "serious": "High",
    "significant": "High",
    
    # Medium synonyms
    "moderate": "Medium",
    "med": "Medium",
    "standard": "Medium",
    
    # Low synonyms
    "minor": "Low",
    "minimal": "Low",
    
    # Informational synonyms
    "information": "Informational",
    "info": "Informational",
    "note": "Informational",
    "fyi": "Informational",
}

# In-memory cache for extracted rubrics
# Key: (name, content_hash), Value: SeverityRubric
_rubric_cache: Dict[Tuple[str, str], "SeverityRubric"] = {}


@dataclass
class SeverityRubric:
    """Severity rubric extracted from a template."""
    allowed_labels: List[str]  # Ordered list of valid severity labels
    definitions: Dict[str, str]  # Label -> definition text
    source: str  # "template" or "default"
    
    def is_valid_label(self, label: str) -> bool:
        """Check if a label is in allowed_labels (case-insensitive)."""
        return label.strip().capitalize() in [lbl.capitalize() for lbl in self.allowed_labels]
    
    def get_canonical_label(self, label: str) -> Optional[str]:
        """Get the canonical form of a label from allowed_labels (case-insensitive match)."""
        label_cap = label.strip().capitalize()
        for allowed in self.allowed_labels:
            if allowed.capitalize() == label_cap:
                return allowed
        return None
    
    def clamp_severity(self, label: str) -> Tuple[str, Optional[str]]:
        """
        Clamp an invalid severity label to a valid one.
        
        Returns:
            (clamped_label, clamp_reason or None)
            
        Strategy:
        1. Try synonym map (deterministic)
        2. Fallback to lowest allowed label
        """
        label = label.strip()
        if not label:
            return self.allowed_labels[-1], "empty_label_clamped_to_lowest"
        
        # Already valid?
        if self.is_valid_label(label):
            return label.capitalize(), None
        
        # Try synonym map (case-insensitive)
        label_lower = label.lower()
        for synonym, canonical in SEVERITY_SYNONYMS.items():
            if label_lower == synonym.lower():
                # Check if canonical is in allowed_labels
                if self.is_valid_label(canonical):
                    return canonical, f"synonym_map:{label}->{canonical}"
        
        # Fallback: use lowest allowed label
        lowest = self.allowed_labels[-1]
        return lowest, f"invalid_label:'{label}' clamped_to_lowest '{lowest}'"


def extract_rubric_from_docx_bytes(name: str, docx_bytes: bytes) -> SeverityRubric:
    """
    Extract severity rubric from DOCX bytes.
    
    Tries multiple extraction strategies in order:
    1. Tables with "Severity" or "Rating" headers
    2. Headings containing "Severity" with label definitions below
    3. Bullets/paragraphs where severity labels appear near "Severity"
    
    Returns:
        SeverityRubric with source="template" if extraction succeeds,
        or source="default" with default labels if extraction fails.
    """
    try:
        # Ensure docx_bytes is bytes (SQLModel sometimes returns str)
        if isinstance(docx_bytes, str):
            docx_bytes = docx_bytes.encode('latin-1')  # Preserve binary data
        
        # Check cache first
        content_hash = hashlib.sha256(docx_bytes).hexdigest()[:16]
        cache_key = (name, content_hash)
        
        if cache_key in _rubric_cache:
            return _rubric_cache[cache_key]
        
        doc = Document(io.BytesIO(docx_bytes))
    except Exception as e:
        log.warning(f"Failed to parse DOCX for rubric extraction: {e}")
        return _default_rubric()
    
    # Try extraction strategies
    rubric = (_extract_from_tables(doc) or 
              _extract_from_headings(doc) or 
              _extract_from_paragraphs(doc))
    
    if rubric:
        rubric.source = "template"
        _rubric_cache[cache_key] = rubric
        
        # Clean up old cache entries (keep only last 10)
        if len(_rubric_cache) > 10:
            oldest_key = next(iter(_rubric_cache))
            del _rubric_cache[oldest_key]
        
        log.info(f"Extracted severity rubric from template '{name}': {rubric.allowed_labels}")
        return rubric
    
    # No rubric found - return default
    log.info(f"No severity rubric found in template '{name}', using defaults")
    return _default_rubric()


def _default_rubric() -> SeverityRubric:
    """Return default severity rubric."""
    return SeverityRubric(
        allowed_labels=DEFAULT_SEVERITY_LABELS,
        definitions={},
        source="default"
    )


def _extract_from_tables(doc: Document) -> Optional[SeverityRubric]:
    """
    Strategy 1: Extract from tables with "Severity" or "Rating" headers.
    
    Looks for tables where:
    - Header row contains "Severity", "Rating", or "Level"
    - First column contains severity labels
    - Second column contains definitions
    """
    for table in doc.tables:
        if len(table.rows) < 2:
            continue
        
        # Check header row for severity-related keywords
        header_text = " ".join(cell.text.strip().lower() for cell in table.rows[0].cells)
        if not any(kw in header_text for kw in ["severity", "rating", "level", "risk"]):
            continue
        
        # Extract labels and definitions from rows
        allowed_labels = []
        definitions = {}
        
        for row in table.rows[1:]:  # Skip header row
            if len(row.cells) < 2:
                continue
            
            cell_text = row.cells[0].text.strip()
            # Check if this cell contains a severity label
            match = re.match(r'^(Critical|High|Medium|Low|Informational)\b', cell_text, re.IGNORECASE)
            if match:
                label = match.group(1).capitalize()
                definition = row.cells[1].text.strip()
                
                if label not in allowed_labels:
                    allowed_labels.append(label)
                if definition and label not in definitions:
                    definitions[label] = definition
        
        if allowed_labels:
            log.info(f"Extracted rubric from table: {allowed_labels}")
            return SeverityRubric(
                allowed_labels=allowed_labels,
                definitions=definitions,
                source="template"
            )
    
    return None


def _extract_from_headings(doc: Document) -> Optional[SeverityRubric]:
    """
    Strategy 2: Extract from headings containing "Severity" with definitions below.
    
    Looks for patterns like:
    - "Severity Definitions" heading
    - Followed by "Critical:" paragraph
    - Followed by definition paragraph(s)
    """
    allowed_labels = []
    definitions = {}
    
    # Find paragraphs with "Severity" in heading style
    severity_section_start = None
    for i, para in enumerate(doc.paragraphs):
        if para.style.name.startswith('Heading') and 'severity' in para.text.lower():
            severity_section_start = i
            break
    
    if severity_section_start is None:
        return None
    
    # Parse following paragraphs for severity labels and definitions
    pattern = re.compile(
        r'^(Critical|High|Medium|Low|Informational)\s*(?:Severity)?[:\-–—]\s*(.+)',
        re.IGNORECASE
    )
    
    for para in doc.paragraphs[severity_section_start + 1:]:
        # Stop at next heading
        if para.style.name.startswith('Heading'):
            break
        
        text = para.text.strip()
        if not text:
            continue
        
        match = pattern.match(text)
        if match:
            label = match.group(1).capitalize()
            definition = match.group(2).strip()
            
            if label not in allowed_labels:
                allowed_labels.append(label)
            if definition and label not in definitions:
                definitions[label] = definition
    
    if allowed_labels:
        log.info(f"Extracted rubric from headings: {allowed_labels}")
        return SeverityRubric(
            allowed_labels=allowed_labels,
            definitions=definitions,
            source="template"
        )
    
    return None


def _extract_from_paragraphs(doc: Document) -> Optional[SeverityRubric]:
    """
    Strategy 3: Extract from bullets/paragraphs where severity labels appear.
    
    Looks for patterns like:
    - "Critical - definition text"
    - "High: definition text"
    """
    allowed_labels = []
    definitions = {}
    
    pattern = re.compile(
        r'\b(Critical|High|Medium|Low|Informational)\s*[\-–—:]\s*(.{10,})',
        re.IGNORECASE
    )
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text or 'severity' not in text.lower():
            continue
        
        for match in pattern.finditer(text):
            label = match.group(1).capitalize()
            definition = match.group(2).strip()
            
            if label not in allowed_labels:
                allowed_labels.append(label)
            if definition and label not in definitions:
                definitions[label] = definition[:200]  # Limit definition length
    
    if allowed_labels:
        log.info(f"Extracted rubric from paragraphs: {allowed_labels}")
        return SeverityRubric(
            allowed_labels=allowed_labels,
            definitions=definitions,
            source="template"
        )
    
    return None


def enforce_severity_label(
    structured: Dict,
    rubric: SeverityRubric,
    notes: str = ""
) -> Dict:
    """
    Enforce that severity field matches rubric allowed_labels.
    
    Clamps invalid labels deterministically using synonym map.
    Never runs regeneration loops - single-pass enforcement.
    
    Adds internal flags for traceability:
    - structured["_severity_flags"] = ["clamped", "source=template", etc.]
    """
    severity_field = None
    severity_value = None
    
    # Find severity field (case-insensitive)
    for key in structured.keys():
        if key.lower() == "severity":
            severity_field = key
            severity_value = structured[key]
            break
    
    if not severity_field:
        # No severity field found - add one with lowest label
        structured["Severity"] = rubric.allowed_labels[-1]
        structured.setdefault("_severity_flags", []).append("missing_severity_added_lowest")
        log.warning(f"No severity field found, added: {rubric.allowed_labels[-1]}")
        return structured
    
    # Clamp if invalid
    if not rubric.is_valid_label(severity_value):
        clamped, reason = rubric.clamp_severity(severity_value)
        structured[severity_field] = clamped
        
        # Add traceability flags
        flags = structured.setdefault("_severity_flags", [])
        if reason:
            flags.append(reason)
        flags.append(f"rubric_source={rubric.source}")  # "template"=DOCX master template, "default"=fallback
        
        log.warning(f"Clamped severity '{severity_value}' -> '{clamped}' (reason: {reason})")
    else:
        # Valid - normalize to canonical form from allowed_labels
        canonical = rubric.get_canonical_label(severity_value)
        if canonical:
            structured[severity_field] = canonical
        # Just add source flag
        structured.setdefault("_severity_flags", []).append(f"rubric_source={rubric.source}")  # "template"=DOCX master template, "default"=fallback
    
    return structured
